import warnings
from dataclasses import replace

import pytest
from pptx import Presentation

from pptato import (
    Bullets,
    Columns,
    Deck,
    Footnote,
    LayoutError,
    LayoutWarning,
    Spacer,
    Stack,
    Table,
    Text,
    TextFit,
    TextStyle,
    Theme,
)
from pptato.layout import Engine


def table(rows=8, **kwargs):
    return Table(["Item", "Value"], [[f"Row {i:03}", str(i)] for i in range(rows)], **kwargs)


def table_nodes(layout):
    return [node for slide in layout.slides for node in slide.walk() if node.kind == "table"]


def height_at(block, size, theme=None, width=888):
    engine = Engine(theme or Theme())
    plain = replace(block, overflow="error", fit=None, style=TextStyle(size=size, paragraph_gap=0))
    return engine.resolve(plain, 0, 0, width, float("inf"), "table").box.height


def fit_at(size, *, policy=TextFit(), block=None):
    block = block or table()
    height = height_at(block, size)
    deck = Deck(size=(960, height + 72))
    deck.add_slide(title="", body=replace(block, overflow="shrink", fit=policy))
    return deck.layout()


def test_default_policy_still_errors_and_never_adds_slides(tmp_path):
    deck = Deck()
    deck.add_slide(title="Too much", body=table(100))
    output = tmp_path / "existing.pptx"
    output.write_bytes(b"unchanged")
    with pytest.raises(LayoutError):
        deck.save(output)
    assert output.read_bytes() == b"unchanged"
    assert len(deck.slides) == 1


def test_normal_shrink_selects_largest_candidate_and_records_decision():
    with warnings.catch_warnings(record=True) as caught:
        layout = fit_at(13.25)
    assert not caught
    fit = table_nodes(layout)[0].table_fit
    assert fit.preferred_size == 15
    assert fit.selected_size == 13.25
    assert fit.status == "normal"
    assert not layout.diagnostics
    assert len(layout.slides) == 1


def test_preferred_size_retained_when_it_fits():
    layout = fit_at(18)
    assert table_nodes(layout)[0].table_fit.selected_size == 15


def test_warning_band_renders_with_structured_and_python_warning():
    with pytest.warns(LayoutWarning, match="below the normal minimum") as caught:
        layout = fit_at(11.25)
    assert len(caught) == 1
    assert len(layout.slides) == 1
    assert table_nodes(layout)[0].table_fit.selected_size == 11.25
    diagnostic = layout.diagnostics[0]
    assert diagnostic.severity == "warning"
    assert diagnostic.code == "readability_warning"
    assert diagnostic.path == "slide/1/body"
    assert diagnostic.required == 12
    assert diagnostic.available == 11.25
    assert layout.to_dict()["diagnostics"][0]["severity"] == "warning"


@pytest.mark.parametrize("boundary", [12, 12.1])
def test_normal_minimum_is_inclusive_even_off_grid(boundary):
    with warnings.catch_warnings(record=True) as caught:
        layout = fit_at(boundary, policy=TextFit(normal_min=boundary))
    assert not caught
    assert table_nodes(layout)[0].table_fit.selected_size == boundary


@pytest.mark.parametrize("boundary", [10, 10.1])
def test_absolute_minimum_is_inclusive_even_off_grid(boundary):
    with pytest.warns(LayoutWarning):
        layout = fit_at(boundary, policy=TextFit(absolute_min=boundary))
    assert table_nodes(layout)[0].table_fit.selected_size == boundary


def test_below_absolute_minimum_errors_without_warning_or_continuation():
    with warnings.catch_warnings(record=True) as caught:
        with pytest.raises(LayoutError) as error:
            fit_at(9.5)
    assert error.value.diagnostic.code == "text_fit_limit"
    assert "absolute minimum 10 pt" in str(error.value)
    assert not caught


def test_no_warning_band_when_thresholds_match():
    with pytest.raises(LayoutError):
        fit_at(11, policy=TextFit(normal_min=12, absolute_min=12))


def test_each_auto_candidate_reallocates_columns_and_remeasures_rows():
    block = Table(
        ["Description", "N"],
        [
            ["Accounts that renewed after a pause in service during the preceding quarter.", "48"],
            ["New accounts that completed the initial source review and onboarding process.", "17"],
        ]
        * 4,
        widths="auto",
    )
    with pytest.warns(LayoutWarning):
        layout = fit_at(11, block=block)
    node = table_nodes(layout)[0]
    selected = node.table_fit.selected_size
    assert 10 <= selected < 12
    expected = Engine(Theme()).resolve(
        replace(block, style=TextStyle(size=selected, paragraph_gap=0)),
        0,
        0,
        888,
        float("inf"),
        "table",
    )
    assert node.column_widths == expected.column_widths
    assert node.row_heights == expected.row_heights
    assert all(cell.style.size == selected for cell in node.children)
    assert all(cell.style.bold for cell in node.children[:2])
    assert not any(cell.style.bold for cell in node.children[2:])


def test_shrink_can_resolve_width_overflow():
    block = Table(["ID"], [["MMMMMMMMMMMM"]], overflow="shrink", fit=TextFit())
    floor = replace(block, overflow="error", fit=None, style=TextStyle(size=11))
    width = Engine(Theme()).intrinsic(floor, "table")[0]
    deck = Deck(size=(width + 72, 540))
    deck.add_slide(title="", body=block)
    with pytest.warns(LayoutWarning):
        layout = deck.layout()
    assert 10 <= table_nodes(layout)[0].table_fit.selected_size < 12


def test_stack_reserves_following_text_before_shrinking_table():
    block = table()
    theme = Theme()
    engine = Engine(theme)
    after = Text("Following commentary")
    after_height = engine.resolve(after, 0, 0, 888, 1000, "after").box.height
    deck = Deck(size=(960, 72 + height_at(block, 13) + theme.gap + after_height))
    deck.add_slide(title="", body=Stack([replace(block, overflow="shrink", fit=TextFit()), after]))
    layout = deck.layout()
    assert table_nodes(layout)[0].table_fit.selected_size == 13
    assert any(node.kind == "text" and node.path.endswith("/1") for node in layout.slides[0].walk())


def test_shrink_in_auto_columns_uses_permitted_content_minimum():
    block = table(5, widths="auto", overflow="shrink", fit=TextFit())
    deck = Deck()
    deck.add_slide(title="Mixed", body=Columns([block, Bullets(["Commentary"])], widths="auto"))
    layout = deck.layout()
    assert table_nodes(layout)[0].table_fit.selected_size == 15


def test_failed_later_slide_does_not_emit_speculative_warnings():
    block = table()
    height = height_at(block, 11)
    deck = Deck(size=(960, height + 72))
    deck.add_slide(title="", body=replace(block, overflow="shrink", fit=TextFit()))
    deck.add_slide(title="", body=Spacer(10000))
    with warnings.catch_warnings(record=True) as caught:
        with pytest.raises(LayoutError):
            deck.layout()
    assert not caught


def test_continuation_preserves_rows_headers_widths_notes_and_source():
    source = table(30, widths="auto", overflow="continue")
    deck = Deck()
    deck.add_slide(title="Inventory", body=source, notes=[Footnote("Source: example inventory.")])
    deck.add_slide(title="Following slide", body=Text("After the table"))
    first, second = deck.layout(), deck.layout()
    assert first == second
    assert len(deck.slides) == 2
    nodes = table_nodes(first)
    assert len(nodes) > 1
    rows = []
    for index, node in enumerate(nodes):
        assert node.column_widths == nodes[0].column_widths
        assert node.children[0].text.paragraphs[0].lines == ("Item",)
        assert node.children[1].text.paragraphs[0].lines == ("Value",)
        rows.extend(cell.text.paragraphs[0].lines[0] for cell in node.children[2::2])
        meta = node.continuation
        assert meta.source_slide == 1
        assert meta.part == index + 1
        assert meta.parts == len(nodes)
        assert meta.row_end > meta.row_start
        assert first.slides[index].continuation == meta
        texts = [
            paragraph.lines
            for n in first.slides[index].walk()
            if n.text
            for paragraph in n.text.paragraphs
        ]
        assert ("Source: example inventory.",) in texts
        expected_title = "Inventory" if index == 0 else "Inventory (continued)"
        assert (expected_title,) in texts
    assert rows == [row[0] for row in source.rows]
    assert first.slides[-1].path == f"slide/{len(first.slides)}"
    assert first.slides[-1].continuation is None
    assert source.overflow == "continue" and len(source.rows) == 30


def test_continuation_retains_original_striping_and_row_paths():
    deck = Deck(size=(960, 270))
    source = table(16, overflow="continue")
    deck.add_slide(title="", body=source)
    nodes = table_nodes(deck.layout())
    for node in nodes:
        for cell in node.children[2:]:
            row_index = int(cell.path.split("/row/")[1].split("/")[0])
            expected = (
                deck.theme.table_alternate_fill if row_index % 2 == 0 else deck.theme.table_fill
            )
            assert cell.fill == expected


def test_continuation_exact_fit_does_not_add_empty_page():
    source = table(5, overflow="continue")
    deck = Deck(size=(960, height_at(source, 15) + 72))
    deck.add_slide(title="", body=source)
    layout = deck.layout()
    assert len(layout.slides) == 1
    assert table_nodes(layout)[0].continuation.row_end == 5


def test_header_only_continuation_produces_one_table():
    deck = Deck()
    deck.add_slide(title="Empty", body=table(0, overflow="continue"))
    layout = deck.layout()
    assert len(layout.slides) == 1
    assert len(table_nodes(layout)[0].row_heights) == 1


def test_too_tall_row_raises_instead_of_looping_or_splitting_cell():
    deck = Deck()
    deck.add_slide(
        title="Too tall", body=Table(["Description"], [["Line\n" * 100]], overflow="continue")
    )
    with pytest.raises(LayoutError) as error:
        deck.layout()
    assert error.value.diagnostic.code == "continuation_row_too_tall"
    assert error.value.diagnostic.path.endswith("/row/1")


def test_continuation_does_not_fix_horizontal_overflow():
    deck = Deck()
    deck.add_slide(title="Too wide", body=Table(["ID"], [["W" * 300]], overflow="continue"))
    with pytest.raises(LayoutError) as error:
        deck.layout()
    assert error.value.diagnostic.axis == "width"


@pytest.mark.parametrize("container", [Stack, Columns])
def test_nested_continuation_is_rejected_explicitly(container):
    deck = Deck()
    deck.add_slide(title="Nested", body=container([table(30, overflow="continue")]))
    with pytest.raises(ValueError, match="requires a Table as the slide body"):
        deck.layout()


@pytest.mark.parametrize(
    "factory",
    [
        lambda: TextFit(preferred=0),
        lambda: TextFit(normal_min=16),
        lambda: TextFit(absolute_min=13),
        lambda: TextFit(absolute_min=float("nan")),
        lambda: table(overflow="shrink"),
        lambda: table(fit=TextFit()),
        lambda: table(overflow="continue", fit=TextFit()),
        lambda: table(overflow="clip"),
    ],
)
def test_invalid_policy_is_rejected(factory):
    with pytest.raises(ValueError):
        factory()


def test_shrink_and_continuation_native_output(tmp_path):
    with pytest.warns(LayoutWarning):
        shrunk = fit_at(11)
    path = tmp_path / "shrunk.pptx"
    shrunk.save(path)
    presentation = Presentation(path)
    native = next(shape.table for shape in presentation.slides[0].shapes if shape.has_table)
    assert all(
        cell.text_frame.paragraphs[0].font.size.pt == 11
        for row in native.rows
        for cell in row.cells
    )
    deck = Deck()
    source = table(30, overflow="continue")
    deck.add_slide(title="Inventory", body=source)
    resolved = deck.layout()
    resolved.save(tmp_path / "continued.pptx")
    presentation = Presentation(tmp_path / "continued.pptx")
    assert len(presentation.slides) == len(resolved.slides)
    actual = []
    for slide, node in zip(presentation.slides, table_nodes(resolved)):
        native = next(shape.table for shape in slide.shapes if shape.has_table)
        assert native.cell(0, 0).text == "Item"
        assert [row.height.pt for row in native.rows] == pytest.approx(node.row_heights, abs=0.001)
        actual.extend(native.cell(row, 0).text for row in range(1, len(native.rows)))
    assert actual == [row[0] for row in source.rows]
