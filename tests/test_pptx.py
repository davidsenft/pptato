import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from pptato import Bullets, Deck, Footnote, Stack, Table, Text, Theme


def test_native_content_and_geometry_round_trip(tmp_path):
    deck = Deck(theme=Theme())
    deck.add_slide(
        title="Native objects",
        body=Stack(
            [
                Bullets(["First\nSecond line", "Another item"]),
                Table(["Metric", "Value"], [["Revenue", "125"], ["Margin", "24%"]]),
            ]
        ),
        notes=[Footnote("Example source.")],
    )
    layout = deck.layout()
    path = tmp_path / "example.pptx"
    layout.save(path)
    prs = Presentation(path)
    assert len(prs.slides) == 1
    assert prs.slide_width.pt == 960
    assert prs.slide_height.pt == 540
    slide = prs.slides[0]
    shapes = {shape.name: shape for shape in slide.shapes}
    for node in layout.slides[0].walk():
        if node.kind not in {"text", "table"}:
            continue
        shape = shapes[node.path]
        for actual, expected in zip(
            (shape.left.pt, shape.top.pt, shape.width.pt, shape.height.pt),
            (node.box.x, node.box.y, node.box.width, node.box.height),
        ):
            assert actual == pytest.approx(expected, abs=0.001)
    bullets = shapes["slide/1/body/0"].text_frame
    assert len(bullets.paragraphs) == 2
    assert bullets.paragraphs[0].text == "First\vSecond line"
    assert len(bullets._txBody.findall(".//" + qn("a:buChar"))) == 2
    tags = [element.tag for element in bullets.paragraphs[0]._p.get_or_add_pPr()]
    assert tags.index(qn("a:lnSpc")) < tags.index(qn("a:buFont"))
    assert tags.index(qn("a:buFont")) < tags.index(qn("a:buChar")) < tags.index(qn("a:defRPr"))
    table = shapes["slide/1/body/1"].table
    assert table.cell(1, 0).text == "Revenue"
    assert table.cell(2, 1).text == "24%"
    assert all(
        cell.text_frame.paragraphs[0].font.name == layout.font_name
        for row in table.rows
        for cell in row.cells
    )


def test_failed_layout_does_not_overwrite_existing_file(tmp_path):
    from pptato import LayoutError

    path = tmp_path / "existing.pptx"
    path.write_bytes(b"existing data")
    deck = Deck()
    deck.add_slide(title="Too much", body=Text("Text " * 10000))
    with pytest.raises(LayoutError):
        deck.save(path)
    assert path.read_bytes() == b"existing data"
