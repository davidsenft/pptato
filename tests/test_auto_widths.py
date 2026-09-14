from dataclasses import replace

import pytest
from pptx import Presentation

from pptato import (
    Bullets,
    Columns,
    ColumnWidth,
    Deck,
    Footnote,
    Heading,
    LayoutError,
    Row,
    Spacer,
    Stack,
    Table,
    Text,
    Theme,
)
from pptato.layout import Engine
from pptato.sizing import WidthProfile, allocate_auto


def layout_for(body, *, size=(960, 540), notes=()):
    deck = Deck(size=size)
    deck.add_slide(title="Automatic widths", body=body, notes=notes)
    return deck.layout()


def body_of(layout):
    return next(node for node in layout.slides[0].walk() if node.path == "slide/1/body")


def example_table():
    return Table(
        ["Description", "N", "%"],
        [
            [
                "Accounts that renewed after a temporary pause in service during the previous quarter.",
                "128",
                "24%",
            ],
            [
                "New accounts that completed onboarding and submitted their first monthly report.",
                "46",
                "18%",
            ],
            [
                "Existing accounts that expanded their subscription to include an additional region.",
                "73",
                "31%",
            ],
        ],
    )


def test_auto_gives_prose_more_width_and_reduces_actual_row_heights():
    table = example_table()
    equal = body_of(layout_for(table))
    auto = body_of(layout_for(replace(table, widths="auto")))
    assert auto.column_widths[0] > equal.column_widths[0]
    assert auto.column_widths[0] > auto.column_widths[1]
    assert auto.box.height < equal.box.height
    assert sum(auto.column_widths) == pytest.approx(auto.box.width)
    assert len(auto.children) == len(equal.children)


def test_long_identifier_gets_its_intrinsic_width():
    text = "ABC-2026-000017928435901"
    table = Table(["ID", "Note"], [[text, "Short note"]], widths="auto")
    node = body_of(layout_for(table, size=(520, 540)))
    cell = node.children[2]
    assert cell.text.paragraphs[0].lines == (text,)
    assert node.column_widths[0] >= node.column_requirements[0].minimum


def test_bold_headers_and_padding_participate_in_minimum():
    theme = Theme(cell_padding=13)
    engine = Engine(theme)
    table = Table(["MMMMMMMMMMMM", "N"], [["a", "1"]], widths="auto")
    expected = (
        engine.measurer.intrinsic((table.headers[0],), replace(theme.table, bold=True))[0] + 26
    )
    assert engine.table_profiles(table)[0].minimum == pytest.approx(expected)


def test_bounds_are_hard_and_unused_space_is_outside_table(tmp_path):
    table = Table(
        ["A", "B"],
        [["One", "Two"]],
        widths="auto",
        bounds=[ColumnWidth(minimum=90, maximum=100), ColumnWidth(maximum=120)],
    )
    layout = layout_for(table)
    node = body_of(layout)
    assert node.column_widths == pytest.approx((100, 120))
    assert node.box.width == pytest.approx(220)
    path = tmp_path / "bounded.pptx"
    layout.save(path)
    native = next(shape for shape in Presentation(path).slides[0].shapes if shape.has_table)
    assert native.width.pt == pytest.approx(220, abs=0.001)


def test_conflicting_content_and_maximum_has_specific_diagnostic():
    table = Table(
        ["Identifier"],
        [["UNBREAKABLE-IDENTIFIER-000142"]],
        widths="auto",
        bounds=[ColumnWidth(maximum=80)],
    )
    with pytest.raises(LayoutError) as error:
        layout_for(table)
    diagnostic = error.value.diagnostic
    assert diagnostic.code == "width_constraints"
    assert diagnostic.path == "slide/1/body/column/0"
    assert diagnostic.required > diagnostic.available == 80
    assert "unbreakable" in diagnostic.suggestion


def test_combined_minima_overflow_reports_whole_allocation():
    table = Table(
        ["A", "B"],
        [["One", "Two"]],
        widths="auto",
        bounds=[ColumnWidth(minimum=500), ColumnWidth(minimum=500)],
    )
    with pytest.raises(LayoutError) as error:
        layout_for(table)
    assert error.value.diagnostic.path == "slide/1/body/columns"
    assert error.value.diagnostic.required == 1000
    assert error.value.diagnostic.available == 888


def test_auto_table_fit_with_footer_where_equal_widths_overflow():
    table = example_table()
    # Use the measured sizes so the test works across our supported font families.
    equal_height = body_of(layout_for(table)).box.height
    auto_height = body_of(layout_for(replace(table, widths="auto"))).box.height
    theme = Theme()
    chrome = 2 * theme.margin + theme.title.size * theme.title.line_height + 2 + theme.title_gap
    note_height = theme.note.size * theme.note.line_height + 2 + theme.footer_gap
    height = chrome + note_height + (equal_height + auto_height) / 2
    notes = [Footnote("Source: synthetic data.")]
    with pytest.raises(LayoutError):
        layout_for(table, size=(960, height), notes=notes)
    node = body_of(layout_for(replace(table, widths="auto"), size=(960, height), notes=notes))
    assert node.box.height == pytest.approx(auto_height)


def test_width_changes_remeasure_auto_table_and_preserve_hard_breaks():
    table = Table(["Summary", "N"], [["First line\nSecond line\n\nLast line", "7"]], widths="auto")
    wide = body_of(layout_for(table))
    narrow = body_of(layout_for(table, size=(300, 540)))
    assert narrow.box.width < wide.box.width
    assert narrow.box.height >= wide.box.height
    assert "" in narrow.children[2].text.paragraphs[0].lines


def test_nested_table_and_commentary_use_same_auto_allocator():
    table = Table(["Identifier", "Count"], [["NORTH-2026-000128", "17"]], widths="auto")
    body = Columns(
        [table, Stack([Heading("Notes"), Bullets(["A longer explanation. " * 5])])],
        widths="auto",
        bounds=[ColumnWidth(minimum=280, maximum=460), ColumnWidth(minimum=220)],
    )
    layout = layout_for(body)
    row = body_of(layout)
    assert 280 <= row.column_widths[0] <= 460
    assert row.column_widths[1] >= 220
    assert sum(row.column_widths) + Theme().gap == pytest.approx(row.box.width)
    for node in layout.slides[0].walk():
        for child in node.children:
            assert child.box.right <= node.box.right + 1e-6
            assert child.box.bottom <= node.box.bottom + 1e-6


@pytest.mark.parametrize("mode", ["equal", [1, 2]])
def test_fixed_width_nested_rows_preserve_their_intrinsic_minimum(mode):
    child = Row([Text("UNBREAKABLE"), Text("Word")], widths=mode, padding=12)
    layout = layout_for(Columns([child, Text("Notes")], widths="auto"))
    assert body_of(layout).children[0].children[0].text.paragraphs[0].lines == ("UNBREAKABLE",)


def test_auto_deterministic_and_snapshot_records_requirements():
    first = layout_for(replace(example_table(), widths="auto"))
    second = layout_for(replace(example_table(), widths="auto"))
    assert first == second
    assert body_of(first).column_requirements[0].preferred > 0
    assert "column_requirements" in str(first.to_dict())


def test_empty_and_header_only_tables_and_empty_auto_row():
    for table in (Table(["A", "B"], [], widths="auto"), Table(["", ""], [["", ""]], widths="auto")):
        assert body_of(layout_for(table)).box.height > 0
    assert body_of(layout_for(Columns([], widths="auto"))).box.height == 0


@pytest.mark.parametrize(
    "factory",
    [
        lambda: ColumnWidth(minimum=-1),
        lambda: ColumnWidth(maximum=0),
        lambda: ColumnWidth(maximum=float("inf")),
        lambda: ColumnWidth(minimum=float("nan")),
        lambda: ColumnWidth(minimum=100, maximum=50),
        lambda: Table(["A"], [], bounds=[ColumnWidth()]),
        lambda: Columns([Text("A")], widths="auto", bounds=[ColumnWidth(), ColumnWidth()]),
        lambda: Table(["A"], [], widths="auto", bounds=[12]),
    ],
)
def test_invalid_bounds_fail_at_construction(factory):
    with pytest.raises(ValueError):
        factory()


def test_allocator_crosses_wrapping_plateaus_and_handles_shared_row_bottleneck():
    profiles = [WidthProfile(20, 80, lambda w: (60 if w < 60 else 20,))] * 2
    widths, _ = allocate_auto(120, profiles, (), "table", table=True)
    assert widths == pytest.approx((60, 60))


def test_allocator_respects_minima_and_caps_for_many_budgets():
    profiles = [
        WidthProfile(20, 100, lambda w: (100 / w,)),
        WidthProfile(40, 160, lambda w: (200 / w,)),
        WidthProfile(10, 70, lambda w: (80 / w,)),
    ]
    bounds = [
        ColumnWidth(maximum=90),
        ColumnWidth(minimum=45, maximum=170),
        ColumnWidth(maximum=100),
    ]
    for width in (75, 100, 200, 359, 500):
        values, requirements = allocate_auto(width, profiles, bounds, "columns")
        assert sum(values) == pytest.approx(min(width, 360))
        for value, requirement in zip(values, requirements):
            assert requirement.minimum - 1e-6 <= value <= requirement.maximum + 1e-6


def test_resolved_auto_columns_and_rows_match_native_pptx(tmp_path):
    layout = layout_for(replace(example_table(), widths="auto"))
    node = body_of(layout)
    path = tmp_path / "auto.pptx"
    layout.save(path)
    shape = next(shape for shape in Presentation(path).slides[0].shapes if shape.has_table)
    assert [column.width.pt for column in shape.table.columns] == pytest.approx(
        node.column_widths, abs=0.001
    )
    assert [row.height.pt for row in shape.table.rows] == pytest.approx(node.row_heights, abs=0.001)
    for cell, expected in zip(
        (cell for row in shape.table.rows for cell in row.cells), node.children
    ):
        assert cell.text == "\v".join(expected.text.paragraphs[0].lines)


def test_explicit_modes_still_work_after_dataclass_replace():
    for block in (example_table(), Columns([Text("A"), Text("B")])):
        changed = replace(block, widths=[1, 2, 1] if isinstance(block, Table) else [1, 2])
        assert body_of(layout_for(changed)).box.width > 0


def test_spacer_and_text_in_auto_columns_stay_finite():
    node = body_of(layout_for(Columns([Spacer(30), Text("Hello")], widths="auto")))
    assert node.box.height >= 30
    assert all(value > 0 for value in node.column_widths)
