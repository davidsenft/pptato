"""Native PowerPoint rendering of an already resolved layout."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from ..layout import Layout, Node


def _text(frame, node: Node, font_name: str) -> None:
    assert node.text is not None and node.style is not None
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = False  # Line breaks come from the measurement pass.
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = Pt(node.padding)
    frame.margin_top = frame.margin_bottom = Pt(node.padding)
    style = node.style
    for index, paragraph in enumerate(node.text.paragraphs):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = "\v".join(paragraph.lines)
        p.line_spacing = Pt(node.text.line_height)
        p.space_before = Pt(0)
        p.space_after = Pt(style.paragraph_gap if index < len(node.text.paragraphs) - 1 else 0)
        p.font.name = font_name
        p.font.size = Pt(style.size)
        p.font.bold = style.bold
        p.font.color.rgb = RGBColor.from_string(style.color)
        props = p._p.get_or_add_pPr()
        if paragraph.bullet:
            props.set("marL", str(int(Pt(node.text.indent))))
            props.set("indent", str(-int(Pt(node.text.indent * 0.75))))
            bullet_font = OxmlElement("a:buFont")
            bullet_font.set("typeface", font_name)
            props.insert_element_before(
                bullet_font,
                "a:buNone",
                "a:buAutoNum",
                "a:buChar",
                "a:buBlip",
                "a:tabLst",
                "a:defRPr",
                "a:extLst",
            )
            bullet = OxmlElement("a:buChar")
            bullet.set("char", "•")
            props.insert_element_before(bullet, "a:tabLst", "a:defRPr", "a:extLst")
        else:
            props.insert_element_before(OxmlElement("a:buNone"), "a:tabLst", "a:defRPr", "a:extLst")


def _render(slide, node: Node, font_name: str) -> None:
    box = node.box
    bounds = tuple(Pt(value) for value in (box.x, box.y, box.width, box.height))
    if node.kind == "text" and node.text and node.text.paragraphs:
        shape = slide.shapes.add_textbox(*bounds)
        shape.name = node.path
        _text(shape.text_frame, node, font_name)
    elif node.kind == "table":
        shape = slide.shapes.add_table(len(node.row_heights), len(node.column_widths), *bounds)
        shape.name = node.path
        table = shape.table
        table.first_row = True
        table.horz_banding = False
        for index, width in enumerate(node.column_widths):
            table.columns[index].width = Pt(width)
        for index, height in enumerate(node.row_heights):
            table.rows[index].height = Pt(height)
        for index, cell_node in enumerate(node.children):
            row, col = divmod(index, len(node.column_widths))
            cell = table.cell(row, col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(cell_node.fill)
            _text(cell.text_frame, cell_node, font_name)
    elif node.kind == "picture":
        shape = slide.shapes.add_picture(BytesIO(node.image_data), *bounds)
        shape.name = node.path
    else:
        for child in node.children:
            _render(slide, child, font_name)


def save(layout: Layout, path: str | Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Pt(layout.size[0])
    presentation.slide_height = Pt(layout.size[1])
    presentation.core_properties.title = "pptato presentation"
    presentation.core_properties.author = "pptato"
    for root in layout.slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _render(slide, root, layout.font_name)
    presentation.save(str(path))
